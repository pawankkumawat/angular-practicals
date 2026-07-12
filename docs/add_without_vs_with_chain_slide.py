"""Add WITHOUT chain vs WITH chain comparison slide to the existing deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX = r"E:\githubrepos\angular-practicals\docs\Angular22-rxResource-Chaining.pptx"

BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
CODE_BG = RGBColor(0x0B, 0x12, 0x20)
ACCENT = RGBColor(0xDD, 0x00, 0x31)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
SOFT_RED = RGBColor(0xFB, 0x71, 0x85)


def set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_multiline(slide, left, top, width, height, lines, size=11, color=LIGHT, font="Consolas", spacing=0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line if line != "" else " "
        # highlight comments / key words lightly
        if line.strip().startswith("//"):
            set_run(run, size=size, color=MUTED, font=font)
        elif "chain(" in line or "WITH CHAIN" in line:
            set_run(run, size=size, bold=True, color=TEAL, font=font)
        elif "WITHOUT CHAIN" in line or "if (!" in line or "return null" in line or "return EMPTY" in line:
            set_run(run, size=size, color=SOFT_RED if "WITHOUT" in line or "if (!" in line or "return null" in line or "EMPTY" in line else LIGHT, font=font)
        else:
            set_run(run, size=size, color=color, font=font)
    return box


def add_card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def move_slide(presentation, old_index, new_index):
    sld_id_lst = presentation.slides._sldIdLst
    slides = list(sld_id_lst)
    el = slides[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


prs = Presentation(PPTX)

insert_after = None
for i, slide in enumerate(prs.slides):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip())
    joined = "\n".join(texts)
    if "What chain Actually Does" in joined:
        insert_after = i
        break
    if insert_after is None and "How chain() propagates status" in joined:
        insert_after = i

slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.28), Inches(0.1), Inches(0.45))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
add_text_box(
    slide, 0.75, 0.22, 12, 0.55,
    "WITHOUT chain vs WITH chain",
    size=24, bold=True, color=WHITE,
)

# LEFT: without chain
add_card(slide, 0.4, 0.9, 6.15, 6.3, fill=CODE_BG)
add_text_box(slide, 0.6, 1.0, 5.8, 0.35, "1. WITHOUT chain", size=16, bold=True, color=SOFT_RED)
add_text_box(
    slide, 0.6, 1.35, 5.8, 0.7,
    "No built-in wait for another resource. You write manual, defensive guards in params and stream.",
    size=11, color=MUTED,
)

without_code = [
    "// WITHOUT CHAIN: Messy, manual synchronization",
    "userResource = rxResource({",
    "  params: () => ({ name: this.currentUsername() }),",
    "  stream: ({ params }) =>",
    "    this.http.get<User>(`/api/users/${params.name}`)",
    "});",
    "",
    "ordersResource = rxResource({",
    "  params: () => {",
    "    // 1. Manually read the value",
    "    const user = this.userResource.value();",
    "",
    "    // 2. If user hasn't loaded, pause",
    "    if (!user) return null;",
    "    return { userId: user.id };",
    "  },",
    "  stream: ({ params }) => {",
    "    // 3. Extra defensive check",
    "    if (!params) return EMPTY;",
    "    return this.http.get<Order[]>(",
    "      `/api/orders?userId=${params.userId}`",
    "    );",
    "  }",
    "});",
]
add_multiline(slide, 0.55, 2.05, 5.85, 5.0, without_code, size=10)

# RIGHT: with chain
add_card(slide, 6.75, 0.9, 6.15, 6.3, fill=CODE_BG)
add_text_box(slide, 6.95, 1.0, 5.8, 0.35, "2. WITH chain", size=16, bold=True, color=TEAL)
add_text_box(
    slide, 6.95, 1.35, 5.8, 0.7,
    "Angular understands the relationship. Boilerplate and manual guard checks disappear.",
    size=11, color=MUTED,
)

with_code = [
    "// WITH CHAIN: Clean, declarative, unified",
    "userResource = rxResource({",
    "  params: () => ({ name: this.currentUsername() }),",
    "  stream: ({ params }) =>",
    "    this.http.get<User>(`/api/users/${params.name}`)",
    "});",
    "",
    "ordersResource = rxResource({",
    "  params: ({ chain }) => {",
    "    // 1. Angular pauses + handles states",
    "    const user = chain(this.userResource);",
    "    return { userId: user.id };",
    "  },",
    "  // 2. No defensive checks needed;",
    "  //    params are ready when this runs",
    "  stream: ({ params }) =>",
    "    this.http.get<Order[]>(",
    "      `/api/orders?userId=${params.userId}`",
    "    )",
    "});",
]
add_multiline(slide, 6.9, 2.05, 5.85, 5.0, with_code, size=10)

new_index = len(prs.slides) - 1
if insert_after is not None:
    move_slide(prs, new_index, insert_after + 1)
    placed = insert_after + 2
else:
    placed = len(prs.slides)

prs.save(PPTX)
print(f"Added slide at position {placed} of {len(prs.slides)}")
print(PPTX)
