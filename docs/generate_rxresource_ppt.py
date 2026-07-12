from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
ACCENT = RGBColor(0xDD, 0x00, 0x31)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
CODE_BG = RGBColor(0x0B, 0x12, 0x20)
TEAL_CARD = RGBColor(0x13, 0x2A, 0x2A)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    font="Calibri",
):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_bullets(slide, left, top, width, height, items, size=18, color=LIGHT, spacing=10):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=color)
    return box


def add_card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_accent_bar(slide, left, top, width=0.12, height=0.55):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    return bar


def title_slide(title, subtitle, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()
    add_text_box(slide, 0.8, 2.2, 11.5, 1.2, title, size=40, bold=True, color=WHITE)
    add_text_box(slide, 0.8, 3.5, 11.5, 1.2, subtitle, size=22, color=MUTED)
    if footer:
        add_text_box(slide, 0.8, 6.7, 11.5, 0.4, footer, size=14, color=MUTED)
    return slide


def content_slide(title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide, 0.6, 0.35)
    add_text_box(slide, 0.9, 0.3, 11.5, 0.7, title, size=28, bold=True, color=WHITE)
    return slide


# 1 Title
title_slide(
    "Chained rxResource in Angular 22",
    "Sequential dependent APIs without callback hell\ngetUser → getBlogById → getCategoryByBlogId",
    "Angular Practicals  |  Callback Hell module",
)

# 2 Agenda
s = content_slide("Agenda")
add_bullets(
    s,
    0.9,
    1.3,
    11,
    5.5,
    [
        "The problem: sequential dependent APIs",
        "resource vs rxResource — which one fits here?",
        "Features of rxResource",
        "Why chaining matters",
        "Our implementation walkthrough",
        "Template binding & key takeaways",
    ],
    size=22,
    spacing=16,
)

# 3 Problem
s = content_slide("The Problem")
add_card(s, 0.8, 1.2, 5.5, 5.4)
add_text_box(s, 1.1, 1.4, 5, 0.5, "Dependent API chain", size=20, bold=True, color=TEAL)
add_bullets(
    s,
    1.1,
    2.1,
    4.9,
    4,
    [
        "getUser() returns a user",
        "getBlogById(user.id) needs the user id",
        "getCategoryByBlogId(blog.postId) needs the blog",
        "Each call depends on the previous result",
    ],
    size=18,
    spacing=12,
)
add_card(s, 6.7, 1.2, 5.8, 5.4)
add_text_box(s, 7.0, 1.4, 5.3, 0.5, "What we want to avoid", size=20, bold=True, color=ACCENT)
add_bullets(
    s,
    7.0,
    2.1,
    5.2,
    4,
    [
        "Nested .subscribe() (callback hell)",
        "Manual loading / error flags",
        "Forgetting to unsubscribe",
        "Reading upstream.value() without status",
    ],
    size=18,
    spacing=12,
)

# 4 Flow
s = content_slide("API Flow")
steps = [
    ("1", "getUser()", "Load current user"),
    ("2", "getBlogById(id)", "Needs user.id"),
    ("3", "getCategoryByBlogId(postId)", "Needs blog.postId"),
]
for i, (num, title, desc) in enumerate(steps):
    left = 0.9 + i * 4.0
    add_card(s, left, 2.2, 3.6, 3.2)
    circle = s.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(left + 1.3), Inches(2.5), Inches(1.0), Inches(1.0)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = ACCENT if i == 0 else TEAL
    circle.line.fill.background()
    add_text_box(
        s, left + 1.3, 2.65, 1.0, 0.7, num, size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_text_box(
        s, left + 0.2, 3.8, 3.2, 0.6, title, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    add_text_box(
        s, left + 0.2, 4.5, 3.2, 0.6, desc, size=15, color=MUTED, align=PP_ALIGN.CENTER
    )
    if i < 2:
        add_text_box(
            s, left + 3.4, 3.3, 0.6, 0.5, "→", size=28, bold=True, color=MUTED, align=PP_ALIGN.CENTER
        )

# 5 resource vs rxResource
s = content_slide("resource vs rxResource — What to use?")
add_card(s, 0.8, 1.2, 5.7, 5.5)
add_text_box(s, 1.1, 1.4, 5.2, 0.5, "resource", size=24, bold=True, color=WHITE)
add_text_box(s, 1.1, 2.0, 5.2, 0.4, "from @angular/core", size=14, color=MUTED)
add_bullets(
    s,
    1.1,
    2.6,
    5.1,
    3.8,
    [
        "loader returns a Promise",
        "Best for fetch(), async/await APIs",
        "One-shot async result per params change",
        "Also supports stream for multi-value signals",
    ],
    size=17,
    spacing=10,
)
add_card(s, 6.9, 1.2, 5.7, 5.5, fill=TEAL_CARD)
add_text_box(s, 7.2, 1.4, 5.2, 0.5, "rxResource  ✓ our choice", size=22, bold=True, color=TEAL)
add_text_box(s, 7.2, 2.0, 5.2, 0.4, "from @angular/core/rxjs-interop", size=14, color=MUTED)
add_bullets(
    s,
    7.2,
    2.6,
    5.1,
    3.8,
    [
        "stream returns an Observable",
        "Best when APIs already return Observables",
        "Fits HttpClient & existing DataService",
        "Same Resource signals: value, status, error…",
    ],
    size=17,
    spacing=10,
)

# 6 Why rxResource here
s = content_slide("Why rxResource in this project?")
add_card(s, 0.8, 1.3, 11.7, 5.3)
add_bullets(
    s,
    1.2,
    1.7,
    11,
    4.5,
    [
        "DataService methods already return Observable<T> (getUser, getBlogById, getCategoryByBlogId)",
        "rxResource accepts those Observables directly via stream — no toPromise / firstValueFrom needed",
        "resource would force wrapping Observables into Promises (extra glue, more error-prone)",
        "httpResource is ideal for raw HttpClient URLs — here we call a service layer, so rxResource fits better",
        "Result: clean bridge from RxJS services → signal-based UI",
    ],
    size=20,
    spacing=14,
)

# 7 Features of rxResource
s = content_slide("Features of rxResource")
features = [
    ("stream", "Loader that returns Observable — emits value(s) into the resource"),
    ("params", "Reactive inputs; when params change, stream re-runs automatically"),
    ("value()", "Signal with the latest resolved data (synchronous read in templates)"),
    ("status / isLoading", "idle | loading | reloading | resolved | error | local"),
    ("error()", "Signal holding the last failure for UI error states"),
    ("reload()", "Re-fetch with the same params without rewriting logic"),
    ("hasValue()", "Type guard + safe check before reading value()"),
    ("defaultValue", "Optional fallback while loading / idle"),
]
for i, (name, desc) in enumerate(features):
    col = i % 2
    row = i // 2
    left = 0.8 + col * 6.2
    top = 1.2 + row * 1.4
    add_card(s, left, top, 5.9, 1.25)
    add_text_box(s, left + 0.25, top + 0.2, 5.4, 0.4, name, size=18, bold=True, color=TEAL)
    add_text_box(s, left + 0.25, top + 0.6, 5.4, 0.5, desc, size=14, color=LIGHT)

# 8 Why chaining
s = content_slide("Why chaining?")
add_card(s, 0.8, 1.2, 11.7, 2.3)
add_text_box(
    s,
    1.1,
    1.4,
    11,
    0.5,
    "Sequential APIs need the previous response before the next call can start",
    size=20,
    bold=True,
    color=WHITE,
)
add_text_box(
    s,
    1.1,
    2.1,
    11,
    1.0,
    "blog needs user.id   ·   category needs blog.postId\n"
    "Chaining expresses that dependency declaratively — no nested subscribe, no manual state machine.",
    size=18,
    color=MUTED,
)
add_card(s, 0.8, 3.8, 5.7, 2.8)
add_text_box(s, 1.1, 4.0, 5.2, 0.4, "Without chain (avoid)", size=18, bold=True, color=ACCENT)
add_bullets(
    s,
    1.1,
    4.5,
    5.2,
    1.8,
    [
        "Read upstream.value() yourself",
        "Downstream goes idle, not loading",
        "Errors / loading not mirrored",
    ],
    size=16,
    spacing=8,
)
add_card(s, 6.9, 3.8, 5.7, 2.8)
add_text_box(s, 7.2, 4.0, 5.2, 0.4, "With chain (preferred)", size=18, bold=True, color=TEAL)
add_bullets(
    s,
    7.2,
    4.5,
    5.2,
    1.8,
    [
        "params: ({ chain }) => chain(user)?.id",
        "Propagates idle / loading / error",
        "Loader runs only when upstream is ready",
    ],
    size=16,
    spacing=8,
)

# 9 How chain works
s = content_slide("How chain() propagates status")
rows = [
    ("Upstream status", "Downstream effect"),
    ("idle", "Also idle — loader does not run"),
    ("loading / reloading", "Enters loading — waits for upstream"),
    ("error", "Enters error — dependency failure"),
    ("resolved / local", "chain() returns the value → params set → stream runs"),
]
for i, (a, b) in enumerate(rows):
    top = 1.3 + i * 1.05
    fill = ACCENT if i == 0 else CARD
    add_card(s, 0.8, top, 5.5, 0.9, fill=fill)
    add_card(s, 6.6, top, 5.9, 0.9, fill=fill)
    add_text_box(s, 1.0, top + 0.25, 5.1, 0.5, a, size=18, bold=(i == 0), color=WHITE)
    add_text_box(s, 6.8, top + 0.25, 5.5, 0.5, b, size=18, bold=(i == 0), color=WHITE)

# 10 Implementation
s = content_slide("Our implementation")
add_card(s, 0.8, 1.15, 11.7, 5.6, fill=CODE_BG)
code = """user = rxResource({
  stream: () => this.service.getUser(),
});

blog = rxResource({
  params: ({ chain }) => chain(this.user)?.id,
  stream: ({ params: userId }) => this.service.getBlogById(userId),
});

category = rxResource({
  params: ({ chain }) => chain(this.blog)?.postId,
  stream: ({ params: postId }) => this.service.getCategoryByBlogId(postId),
});"""
box = s.shapes.add_textbox(Inches(1.1), Inches(1.4), Inches(11.2), Inches(5.2))
tf = box.text_frame
tf.word_wrap = True
for i, line in enumerate(code.split("\n")):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    run = p.add_run()
    run.text = line
    highlight = "rxResource" in line or "chain" in line
    set_run(run, size=16, color=TEAL if highlight else LIGHT, font="Consolas")

# 11 Runtime
s = content_slide("What happens at runtime")
steps2 = [
    ("user", "stream runs immediately → status: loading → resolved with User"),
    ("blog", "chain(user) waits until user is resolved → then calls getBlogById(user.id)"),
    (
        "category",
        "chain(blog) waits until blog is resolved → then calls getCategoryByBlogId(postId)",
    ),
    ("UI", "Each resource exposes isLoading / value / error — template reacts automatically"),
]
for i, (title, desc) in enumerate(steps2):
    top = 1.25 + i * 1.35
    add_card(s, 0.8, top, 11.7, 1.2)
    add_text_box(s, 1.1, top + 0.2, 2.2, 0.7, title, size=20, bold=True, color=TEAL)
    add_text_box(s, 3.5, top + 0.3, 8.6, 0.7, desc, size=17, color=LIGHT)

# 12 Template
s = content_slide("Template binding")
add_card(s, 0.8, 1.2, 11.7, 5.5, fill=CODE_BG)
tmpl = """@if (user.isLoading()) {
  <p>Loading user...</p>
} @else if (user.hasValue()) {
  <pre>{{ user.value() | json }}</pre>
}

@if (blog.isLoading()) { ... }
@if (category.hasValue()) { ... }
@if (category.error(); as error) {
  <p>Error: {{ error.message }}</p>
}"""
box = s.shapes.add_textbox(Inches(1.1), Inches(1.5), Inches(11.2), Inches(5))
tf = box.text_frame
tf.word_wrap = True
for i, line in enumerate(tmpl.split("\n")):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    run = p.add_run()
    run.text = line
    set_run(run, size=17, color=LIGHT, font="Consolas")

# 13 Choosing API
s = content_slide("Choosing the right API")
table = [
    ("Situation", "Use"),
    ("Service / API returns Observable", "rxResource"),
    ("Promise / fetch / async function", "resource"),
    ("Direct HTTP URL via HttpClient", "httpResource"),
    ("One resource depends on another", "params + chain(...)"),
    ("Derive sync value from a resource", "computed (not chain)"),
]
for i, (a, b) in enumerate(table):
    top = 1.25 + i * 0.95
    fill = ACCENT if i == 0 else CARD
    add_card(s, 0.8, top, 5.8, 0.85, fill=fill)
    add_card(s, 6.9, top, 5.6, 0.85, fill=fill)
    add_text_box(s, 1.05, top + 0.22, 5.3, 0.5, a, size=17, bold=(i == 0), color=WHITE)
    add_text_box(s, 7.15, top + 0.22, 5.1, 0.5, b, size=17, bold=(i == 0), color=WHITE)

# 14 Takeaways
s = content_slide("Key takeaways")
add_bullets(
    s,
    0.9,
    1.4,
    11.5,
    5.5,
    [
        "We chose rxResource because DataService already exposes Observables",
        "resource is for Promises; httpResource is for declarative HTTP URLs",
        "Chaining models real-world dependent APIs without nested callbacks",
        "chain() correctly mirrors upstream loading and error states",
        "Pass chain(...)?.field directly as params — don’t wrap undefined in an object",
        "Templates stay simple: isLoading(), hasValue(), value(), error()",
    ],
    size=22,
    spacing=16,
)

# 15 End
title_slide(
    "Questions?",
    "Angular 22  ·  rxResource + chain()\nCallback Hell module → /cbh",
    "Docs: angular.dev/guide/signals/resource",
)

out = r"E:\githubrepos\angular-practicals\docs\Angular22-rxResource-Chaining.pptx"
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print(out)
