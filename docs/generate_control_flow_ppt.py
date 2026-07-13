from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
ACCENT = RGBColor(0xDD, 0x00, 0x31)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
CODE_BG = RGBColor(0x0B, 0x12, 0x20)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    font="Calibri",
):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_bullets(slide, left, top, width, height, items, size=18, color=LIGHT, spacing=10):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=color)
    return box


def add_card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_accent_bar(slide, left, top, width=0.12, height=0.55):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    return bar


def title_slide(title, subtitle, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()
    add_text_box(slide, 0.8, 2.2, 11.5, 1.2, title, size=40, bold=True, color=WHITE)
    add_text_box(slide, 0.8, 3.5, 11.5, 1.2, subtitle, size=22, color=MUTED)
    if footer:
        add_text_box(slide, 0.8, 6.7, 11.5, 0.4, footer, size=14, color=MUTED)
    return slide


def content_slide(title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_accent_bar(slide, 0.6, 0.35)
    add_text_box(slide, 0.9, 0.3, 11.5, 0.7, title, size=28, bold=True, color=WHITE)
    return slide


def add_code_block(slide, left, top, width, height, code, size=14):
    add_card(slide, left, top, width, height, fill=CODE_BG)
    box = slide.shapes.add_textbox(
        Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.3)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line if line else " "
        highlight = any(
            k in line
            for k in ("@if", "@for", "@else", "ng-container", "*ngIf", "*ngFor")
        )
        set_run(run, size=size, color=TEAL if highlight else LIGHT, font="Consolas")


# ── Slide 1: Title ──────────────────────────────────────────────
title_slide(
    "Angular Control Flow Statements",
    "@if  ·  @for  ·  @switch\nBuilt-in template syntax that replaces *ngIf, *ngFor, and ng-container workarounds",
    "Angular Practicals  |  Ng Container module → /ngc/example",
)

# ── Slide 2: When introduced ────────────────────────────────────
s = content_slide("When was control flow introduced?")
add_card(s, 0.8, 1.2, 5.7, 5.5)
add_text_box(s, 1.1, 1.5, 5.2, 0.5, "Angular 17", size=28, bold=True, color=TEAL)
add_text_box(s, 1.1, 2.2, 5.2, 0.4, "November 2023", size=18, color=MUTED)
add_bullets(
    s,
    1.1,
    2.9,
    5.1,
    3.5,
    [
        "Shipped as a developer preview in Angular 17",
        "New built-in blocks: @if, @for, @switch, @defer",
        "No CommonModule import needed for these blocks",
        "Became the recommended template syntax going forward",
    ],
    size=17,
    spacing=12,
)

add_card(s, 6.9, 1.2, 5.7, 5.5)
add_text_box(s, 7.2, 1.5, 5.2, 0.5, "What you get", size=22, bold=True, color=WHITE)
add_bullets(
    s,
    7.2,
    2.3,
    5.1,
    4.0,
    [
        "@if / @else if / @else — conditional rendering",
        "@for (item of items; track item.id) — loops with mandatory track",
        "@switch / @case / @default — multi-branch UI",
        "Clearer templates, better type narrowing, fewer microsyntax surprises",
    ],
    size=17,
    spacing=12,
)

# ── Slide 3: How it fixes ng-container ──────────────────────────
s = content_slide("How control flow fixes the ng-container problem")
add_text_box(
    s,
    0.9,
    1.0,
    11.5,
    0.5,
    "From ng-container-example.component.html — filter rows where number < 400",
    size=16,
    color=MUTED,
)

# Before card
add_card(s, 0.8, 1.55, 5.9, 5.2)
add_text_box(s, 1.05, 1.7, 5.4, 0.4, "Before — *ngFor + ng-container *ngIf", size=16, bold=True, color=ACCENT)
before = """<tr *ngFor="let item of data">
  <ng-container *ngIf="item.number<400">
    <td>{{item.name}}</td>
    <td>{{item.number}}</td>
    <td>{{item.value}}</td>
  </ng-container>
</tr>"""
add_code_block(s, 1.0, 2.25, 5.5, 2.5, before, size=13)
add_bullets(
    s,
    1.05,
    4.9,
    5.4,
    1.6,
    [
        "Can't put *ngFor and *ngIf on the same <tr>",
        "ng-container is a workaround with no DOM element",
        "Still leaves an empty <tr> when the condition fails",
    ],
    size=14,
    spacing=6,
)

# After card
add_card(s, 6.95, 1.55, 5.7, 5.2)
add_text_box(s, 7.2, 1.7, 5.2, 0.4, "After — @for + @if (no ng-container)", size=16, bold=True, color=TEAL)
after = """@for (item of data; track item) {
  @if (item.number < 400) {
    <tr>
      <td>{{item.name}}</td>
      <td>{{item.number}}</td>
      <td>{{item.value}}</td>
    </tr>
  }
}"""
add_code_block(s, 7.15, 2.25, 5.3, 2.5, after, size=13)
add_bullets(
    s,
    7.2,
    4.9,
    5.2,
    1.6,
    [
        "Nest @if inside @for freely — no microsyntax clash",
        "No ng-container needed as a grouping host",
        "<tr> is created only when the condition is true",
    ],
    size=14,
    spacing=6,
)

out = r"E:\githubrepos\angular-practicals\docs\Angular-Control-Flow-Statements.pptx"
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print(out)
