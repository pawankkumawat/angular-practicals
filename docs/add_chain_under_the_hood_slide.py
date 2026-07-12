"""Add one slide to the existing (user-modified) deck without regenerating it."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX = r"E:\githubrepos\angular-practicals\docs\Angular22-rxResource-Chaining.pptx"

BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1E, 0x29, 0x3B)
ACCENT = RGBColor(0xDD, 0x00, 0x31)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)


def set_run(run, size=18, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def move_slide(presentation, old_index, new_index):
    sld_id_lst = presentation.slides._sldIdLst
    slides = list(sld_id_lst)
    el = slides[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


prs = Presentation(PPTX)

# Find insertion point: after "How chain() propagates status" if present, else after "Why chaining?"
insert_after = None
for i, slide in enumerate(prs.slides):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip())
    joined = "\n".join(texts)
    if "How chain() propagates status" in joined:
        insert_after = i
        break
    if insert_after is None and "Why chaining?" in joined:
        insert_after = i

slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

# Accent bar + title
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.35), Inches(0.12), Inches(0.55))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
add_text_box(
    slide, 0.9, 0.3, 11.5, 0.7,
    "What chain Actually Does (Under the Hood)",
    size=26, bold=True, color=WHITE,
)

# Intro
add_card(slide, 0.8, 1.1, 11.7, 1.15)
add_text_box(
    slide, 1.1, 1.3, 11.2, 0.8,
    "When you pass a parent resource into the chain() function, it automatically handles three complex tasks for you:",
    size=17, color=LIGHT,
)

points = [
    (
        "1. Signals Dependency Tracking",
        "It reads the value of the parent resource inside a reactive context. If the parent's value changes later, the downstream resource automatically re-runs.",
    ),
    (
        "2. State Mirroring",
        "If the parent resource is currently loading or has an error, chain() intercepts the flow and instantly forces the downstream resource to mirror that exact same loading or error state.",
    ),
    (
        "3. Execution Guarding",
        "It pauses the downstream resource's loader. The downstream loader or stream will not execute until the parent resource successfully resolves with a valid value.",
    ),
]

for i, (title, body) in enumerate(points):
    top = 2.5 + i * 1.55
    add_card(slide, 0.8, top, 11.7, 1.4)
    add_text_box(slide, 1.1, top + 0.18, 11.2, 0.4, title, size=18, bold=True, color=TEAL)
    add_text_box(slide, 1.1, top + 0.6, 11.2, 0.7, body, size=15, color=LIGHT)

# Move new slide (currently last) to right after the chaining slides
new_index = len(prs.slides) - 1
if insert_after is not None:
    move_slide(prs, new_index, insert_after + 1)
    placed = insert_after + 2  # 1-based for humans
else:
    placed = len(prs.slides)

prs.save(PPTX)
print(f"Added slide at position {placed} of {len(prs.slides)}")
print(PPTX)
